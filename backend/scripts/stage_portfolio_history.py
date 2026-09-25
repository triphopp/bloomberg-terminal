"""Stage the audited 2024-25 Excel history without changing portfolio balances.

The source workbook mixes purchases, sales and later mark prices. Market OHLC
only checks whether a recorded price was possible on a date; it does not prove
execution or funding. Review rows therefore stay outside trades/cash_ledger.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import sqlite3
import uuid
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = Path.home() / "Downloads" / "Portfolio Performance.xlsx"
DEFAULT_AUDIT = ROOT / "backups" / "portfolio-history-20260926"
DEFAULT_DB = ROOT / "portfolio.db"


def _num(value: str | None) -> float | None:
    if value is None or str(value).strip() == "":
        return None
    return float(value)


def _rows(audit_dir: Path, source_hash: str) -> list[tuple]:
    rows: list[tuple] = []
    with (audit_dir / "price_audit.csv").open(encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            row_id = str(uuid.uuid5(uuid.NAMESPACE_URL, f"{source_hash}|{row['sheet']}|{row['row']}|TRADE"))
            rows.append((
                row_id, source_hash, "TRADE", row["sheet"], int(row["row"]), row["account"],
                row["date"], row["date_exit"] or None, row["symbol"],
                _num(row["entry"]), _num(row["exit"]), _num(row["qty"]),
                _num(row["amount"]), _num(row["pnl"]), None, None,
                row["market_day"] or None, _num(row["market_low"]), _num(row["market_high"]),
                row["entry_check"], row["exit_check"], row["review"], None, None,
                "Market OHLC checks price plausibility only; Excel is a self-kept ledger, not a broker confirmation.",
            ))
    with (audit_dir / "cash_reconciliation.csv").open(encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            row_id = str(uuid.uuid5(uuid.NAMESPACE_URL, f"{source_hash}|Income&expenses|{row['row']}|CASH"))
            rows.append((
                row_id, source_hash, "CASH", "Income&expenses", int(row["row"]), row["account"],
                row["date"], None, None, None, None, None, None, None,
                _num(row["income"]), _num(row["investment"]), None, None, None, None, None,
                row["status"], row["db_id"] or None, row["status"],
                "Compare date, account and signed investment with existing cash_ledger; income can have different meaning.",
            ))
    return rows


def _verify_source_rows(source: Path, rows: list[tuple]) -> None:
    workbook = openpyxl.load_workbook(source, data_only=True)
    for row in rows:
        sheet = workbook[row[3]]
        n = row[4]
        recorded_date = sheet.cell(n, 1).value
        if recorded_date is None or recorded_date.strftime("%Y-%m-%d") != row[6]:
            raise ValueError(f"Source date mismatch at {row[3]}!A{n}")
        if row[2] == "TRADE":
            cells = ((3, row[8]), (5, row[9]), (6, row[10]), (9, row[11]),
                     (10, row[12]), (11, row[13]))
        else:
            cells = ((2, row[14]), (3, row[15]))
        for column, expected in cells:
            actual = sheet.cell(n, column).value
            if isinstance(expected, (int, float)):
                try:
                    numeric_actual = float(actual)
                except (TypeError, ValueError):
                    numeric_actual = None
                if numeric_actual is not None and abs(numeric_actual - expected) <= max(1e-6, abs(expected) * 1e-9):
                    continue
            elif (str(actual).strip().upper() if actual is not None else None) == expected:
                continue
            raise ValueError(f"Source value mismatch at {row[3]}!{sheet.cell(n, column).coordinate}")


def _lot_matches(audit_dir: Path, source_hash: str) -> list[tuple]:
    matches = []
    sheets = {"finansia": "Finansia ", "dime": "Dime", "innovestx": "InnovestX "}
    with (audit_dir / "lot_matches.csv").open(encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            row_id = str(uuid.uuid5(
                uuid.NAMESPACE_URL,
                f"{source_hash}|{sheets[row['account']]}|{row['sale_row']}|TRADE",
            ))
            matches.append((row["status"], row["buy_rows"], _num(row["unmatched_qty"]),
                            _num(row["weighted_source_cost"]), row_id))
    return matches


def _create_table(conn: sqlite3.Connection) -> None:
    conn.execute("""
        CREATE TABLE IF NOT EXISTS portfolio_history_review (
            id TEXT PRIMARY KEY,
            source_sha256 TEXT NOT NULL,
            record_type TEXT NOT NULL CHECK(record_type IN ('TRADE', 'CASH')),
            source_sheet TEXT NOT NULL,
            source_row INTEGER NOT NULL,
            account_id TEXT NOT NULL REFERENCES portfolio_accounts(id),
            recorded_date TEXT NOT NULL,
            recorded_exit_date TEXT,
            symbol TEXT,
            price_entry REAL,
            price_exit REAL,
            quantity REAL,
            amount REAL,
            pnl_amount REAL,
            cash_income REAL,
            cash_investment REAL,
            market_date TEXT,
            market_low REAL,
            market_high REAL,
            entry_price_check TEXT,
            exit_price_check TEXT,
            review_status TEXT NOT NULL,
            matched_cash_id TEXT,
            match_status TEXT,
            method_note TEXT NOT NULL,
            lot_match_status TEXT,
            matched_source_rows TEXT,
            unmatched_quantity REAL,
            source_weighted_cost REAL,
            review_decision TEXT,
            review_note TEXT,
            created_at TEXT NOT NULL DEFAULT (strftime('%Y-%m-%d %H:%M:%f','now')),
            UNIQUE(source_sha256, record_type, source_sheet, source_row)
        )
    """)
    columns = {row[1] for row in conn.execute("PRAGMA table_info(portfolio_history_review)")}
    for name, kind in (("lot_match_status", "TEXT"), ("matched_source_rows", "TEXT"),
                       ("unmatched_quantity", "REAL"), ("source_weighted_cost", "REAL")):
        if name not in columns:
            conn.execute(f"ALTER TABLE portfolio_history_review ADD COLUMN {name} {kind}")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_portfolio_history_review_account_date ON portfolio_history_review(account_id, recorded_date)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_portfolio_history_review_status ON portfolio_history_review(review_status)")


def stage(source: Path, audit_dir: Path, db: Path, apply: bool) -> dict:
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    summary = json.loads((audit_dir / "summary.json").read_text(encoding="utf-8"))
    if digest != summary["source_sha256"]:
        raise ValueError("Workbook differs from the audited source; rerun price and cash audits")
    rows = _rows(audit_dir, digest)
    matches = _lot_matches(audit_dir, digest)
    cash_count = sum(row[2] == "CASH" for row in rows)
    sale_count = sum(row[21] == "sale_candidate" for row in rows)
    if (len(rows) != summary["rows"] + cash_count
            or len({row[0] for row in rows}) != len(rows)
            or len(matches) != sale_count):
        raise ValueError("Audit row count or stable IDs do not match")
    _verify_source_rows(source, rows)
    result = {"source_sha256": digest, "review_rows": len(rows), "statuses": dict(Counter(row[21] for row in rows))}
    if not apply:
        return result

    backup = audit_dir / f"portfolio-pre-history-review-{datetime.now(timezone.utc):%Y%m%d-%H%M%S-%f}.db"
    with sqlite3.connect(f"file:{db}?mode=ro", uri=True) as source_conn:
        with sqlite3.connect(backup) as backup_conn:
            source_conn.backup(backup_conn)
            if backup_conn.execute("PRAGMA integrity_check").fetchone()[0] != "ok":
                raise RuntimeError("SQLite backup failed integrity_check")
    result["backup"] = str(backup)
    with sqlite3.connect(db) as conn:
        conn.execute("PRAGMA foreign_keys=ON")
        before = {table: conn.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]
                  for table in ("trades", "cash_ledger", "dividends", "portfolio_nav_snapshots")}
        _create_table(conn)
        fields = """id,source_sha256,record_type,source_sheet,source_row,account_id,
            recorded_date,recorded_exit_date,symbol,price_entry,price_exit,quantity,
            amount,pnl_amount,cash_income,cash_investment,market_date,market_low,
            market_high,entry_price_check,exit_price_check,review_status,matched_cash_id,
            match_status,method_note"""
        count_before = conn.execute("SELECT COUNT(*) FROM portfolio_history_review WHERE source_sha256=?", (digest,)).fetchone()[0]
        conn.executemany(
            f"INSERT OR IGNORE INTO portfolio_history_review ({fields}) VALUES ({','.join('?' * 25)})",
            rows,
        )
        conn.executemany(
            """UPDATE portfolio_history_review
               SET lot_match_status=?, matched_source_rows=?, unmatched_quantity=?, source_weighted_cost=?
               WHERE id=? AND review_status='sale_candidate'""",
            matches,
        )
        count_after = conn.execute("SELECT COUNT(*) FROM portfolio_history_review WHERE source_sha256=?", (digest,)).fetchone()[0]
        after = {table: conn.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0] for table in before}
        if after != before or count_after != len(rows):
            raise RuntimeError("Portfolio counts changed or staged rows are incomplete")
        result.update(inserted=count_after - count_before, live_portfolio_counts_unchanged=True)
    with sqlite3.connect(f"file:{db}?mode=ro", uri=True) as check:
        if check.execute("PRAGMA integrity_check").fetchone()[0] != "ok":
            raise RuntimeError("Live database failed integrity_check")
    return result


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--audit-dir", type=Path, default=DEFAULT_AUDIT)
    parser.add_argument("--db", type=Path, default=DEFAULT_DB)
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()
    print(json.dumps(stage(args.source, args.audit_dir, args.db, args.apply), ensure_ascii=False, indent=2))
