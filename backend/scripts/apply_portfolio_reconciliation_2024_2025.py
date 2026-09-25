"""Apply a narrowly verified subset of the 2024–25 reconciliation workbook.

The workbooks are evidence, not instructions. All 165 proposed closed trades get
a review decision. Only the explicitly audited source rows below enter trades.
Run without --apply for a dry run; --apply makes a SQLite online backup first.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sqlite3
import uuid
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
RECON = Path(r"C:\Users\markereversey\Downloads\Portfolio Reconciliation 2024-2025.xlsx")
ORIGINAL = Path(r"C:\Users\markereversey\Downloads\Portfolio Performance.xlsx")
DB = ROOT / "backend" / "portfolio.db"
BACKUPS = ROOT / "backend" / "backups" / "portfolio-reconciliation"

# Each row has an actual BUY lot in the source workbook, an in-range exit
# price, a plausible source P&L formula, and no known live trade duplicate.
# Other grade-A rows have date, cost, instrument, or duplicate ambiguities.
APPROVED = frozenset({
    "Dime!44", "Dime!45", "Dime!49", "Dime!52", "Dime!58",
    "Dime!64", "Dime!83", "Dime!90", "Dime!96", "Dime!106",
    "Dime!131", "Dime!139", "Dime!153",
    "Finansia!20", "Finansia!24", "Finansia!29", "Finansia!42",
    "Finansia!92", "Finansia!112", "Finansia!123",
})

EXCEPTION_REASON = {
    "Dime!46": "Exit price is below independent daily low.",
    "Dime!55": "SGOV exit price exceeds independent daily high.",
    "Dime!76": "Recorded exit is a Saturday; execution date not independently proven.",
    "Dime!77": "Recorded exit is a Saturday; entry cost differs from cited buys.",
    "Dime!86": "Reconciled 2024 entry date conflicts with 2025 buy of 2 shares.",
    "Dime!101": "Cited buy date falls on a Saturday and cost allocation is unclear.",
    "Dime!113": "Cited buy price exceeds the independent daily high.",
    "Dime!120": "Cited buy date falls on a Saturday.",
    "Dime!121": "Cited buy date falls on a Saturday; mixed-lot cost needs review.",
    "Dime!128": "Cited buy date falls on a Saturday; source cost differs from buy.",
    "Dime!132": "Cited buy date falls on a Saturday.",
    "Dime!155": "GOLD is a commodity in the sheet; ticker price check is unavailable.",
    "Dime!156": "GOLD is a commodity in the sheet; ticker price check is unavailable.",
    "Dime!157": "Reconciled entry predates the exact-sized buy by two months.",
    "Finansia!36": "Source sell cost 2.54 differs from cited buy price 2.58.",
    "Finansia!58": "Reconciled entry date precedes the cited buy date.",
    "Finansia!61": "Cited buy lots imply a different weighted cost.",
    "Finansia!65": "Reconciled entry date precedes the cited buy date.",
    "Finansia!77": "Exit price 1.83 exceeds independent daily high 1.82.",
    "Finansia!115": "Reconciled entry date precedes the cited buy date.",
}
SHEETS = {"Dime": "Dime", "Finansia": "Finansia ", "InnovestX": "InnovestX "}


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def iso(value) -> str:
    return value.date().isoformat()


def price_in_range(price: float, low: float, high: float) -> bool:
    # Market data is floats; one-tick or material discrepancies are rejected.
    return low - 0.0001 <= price <= high + 0.0001


def analyze(conn: sqlite3.Connection, recon_path: Path, original_path: Path):
    recon = openpyxl.load_workbook(recon_path, read_only=True, data_only=True)
    original = openpyxl.load_workbook(original_path, read_only=True, data_only=True)
    prices = {r[1]: r for r in recon["Trades_Reconciled"].iter_rows(min_row=2, values_only=True)}
    recon_hash, original_hash = digest(recon_path), digest(original_path)
    decisions, pending, capacity = [], [], defaultdict(float)
    seen_refs = set()
    for row in recon["Import_2024_2025"].iter_rows(min_row=2, values_only=True):
        entry, exit_, symbol, sector, buy, sell, _, _, qty, _, pnl, wl, pct, account, confidence, ref, note = row
        sheet_label, row_num = ref.split("!")
        source_sheet = SHEETS[sheet_label]
        source = [c.value for c in original[source_sheet][int(row_num)]]
        staged = conn.execute(
            "SELECT id FROM portfolio_history_review WHERE source_sheet=? AND source_row=? AND record_type='TRADE'",
            (source_sheet, int(row_num)),
        ).fetchone()
        if not staged or ref in seen_refs:
            raise ValueError(f"Source review row missing or duplicated: {ref}")
        seen_refs.add(ref)
        source_matches = (source[2] == symbol and abs(float(source[8]) - float(qty)) <= 1e-6
                          and abs(float(source[5]) - float(sell)) <= 1e-5
                          and abs(float(source[10]) - float(pnl)) <= 0.02)
        if ref in APPROVED and not source_matches:
            raise ValueError(f"Approved source row differs from reconciliation: {ref}")
        if iso(exit_)[:4] not in ("2024", "2025") or not entry or iso(entry) > iso(exit_):
            raise ValueError(f"Invalid historical dates: {ref}")
        market = prices[ref]
        reason = EXCEPTION_REASON.get(ref)
        if not source_matches:
            reason = "Reconciliation figures differ from original source row; " + (reason or "manual review required.")
        if ref in APPROVED:
            if confidence != "A" or account not in ("dime", "finansia"):
                raise ValueError(f"Approved row lost grade/account: {ref}")
            if not market[16] or not price_in_range(float(sell), float(market[16]), float(market[17])):
                raise ValueError(f"Approved exit price outside market day: {ref}")
            if abs((float(sell) - float(buy)) * float(qty) - float(pnl)) > max(0.05, float(qty) * 0.0001):
                raise ValueError(f"Approved P&L arithmetic differs: {ref}")
            refs = re.findall(r"(?:Dime|Finansia)!\d+", note or "")
            if not refs or "RECON" in (note or ""):
                raise ValueError(f"Approved sale lacks actual buy rows: {ref}")
            cited_prices = []
            for buy_ref in refs:
                sh, n = buy_ref.split("!")
                buy_row = [c.value for c in original[SHEETS[sh]][int(n)]]
                if sh != sheet_label or buy_row[2] != symbol or buy_row[0] > exit_ or abs(float(buy_row[4]) - float(buy_row[5])) > 1e-6 or abs(float(buy_row[10] or 0)) > 1e-5:
                    raise ValueError(f"Cited buy invalid for {ref}: {buy_ref}")
                buy_review = conn.execute(
                    "SELECT entry_price_check FROM portfolio_history_review WHERE source_sheet=? AND source_row=? AND record_type='TRADE'",
                    (SHEETS[sh], int(n)),
                ).fetchone()
                if not buy_review or buy_review[0] != "in_range":
                    raise ValueError(f"Cited buy price unconfirmed for {ref}: {buy_ref}")
                cited_prices.append(float(buy_row[4]))
                capacity[buy_ref] = float(buy_row[8])
            if float(buy) < min(cited_prices) * 0.985 or float(buy) > max(cited_prices) * 1.015:
                raise ValueError(f"Approved entry cost outside cited buy range: {ref}")
            reason = "Cited source buys and sale match; exit price inside daily range; gross P&L verified."
            pending.append((row, refs))
            decision = "IMPORTED"
        else:
            reason = reason or ("Reconstructed buy/date or missing lot evidence." if confidence == "C" else "Cost, date, or price evidence needs review.")
            decision = "REVIEW_REQUIRED"
        decisions.append((staged[0], ref, decision, reason, row))
    if len(decisions) != 165 or {r[0][15] for r in pending} != APPROVED:
        raise ValueError("Workbook row count or approved source references changed")
    # For each symbol/account, imported sales must fit all cited buys; no
    # reconstructed opening inventory is permitted in this subset.
    sold = defaultdict(float)
    for row, _ in pending:
        sold[(row[13], row[2])] += float(row[8])
    buys = defaultdict(float)
    for ref, qty in capacity.items():
        sh, n = ref.split("!")
        src = [c.value for c in original[SHEETS[sh]][int(n)]]
        buys[("dime" if sh == "Dime" else "finansia", src[2])] += qty
    for key, qty in sold.items():
        if qty > buys[key] + 1e-5:
            raise ValueError(f"Insufficient recorded buys: {key} {qty} > {buys[key]}")
    return decisions, pending, recon_hash, original_hash


def run(recon_path: Path, original_path: Path, db_path: Path, apply: bool):
    with sqlite3.connect(db_path) as conn:
        conn.row_factory = sqlite3.Row
        decisions, pending, recon_hash, original_hash = analyze(conn, recon_path, original_path)
        workbook = openpyxl.load_workbook(recon_path, read_only=True, data_only=True)
        missing_buys = defaultdict(list)
        for candidate in workbook["Missing_Buys"].iter_rows(min_row=2, values_only=True):
            account, symbol, evidence_date, proposed_date, qty, price, gross, sale_ref, method = candidate
            missing_buys[sale_ref].append({
                "account_id": account, "symbol": symbol,
                "evidence_date": iso(evidence_date) if evidence_date else None,
                "proposed_date": iso(proposed_date) if proposed_date else None,
                "quantity": qty, "price": price,
                "gross": gross, "dating_method": method,
            })
        current = conn.execute("SELECT COUNT(*) FROM trades").fetchone()[0]
        proposal_refs = {d[1] for d in decisions}
        result = {"workbook_sha256": recon_hash, "original_sha256": original_hash,
                  "approved": len(pending), "review_required": len(decisions)-len(pending),
                  "missing_buy_proposals_tracked": sum(len(missing_buys.get(ref, [])) for ref in proposal_refs),
                  "missing_buy_proposals_out_of_scope": sum(len(rows) for ref, rows in missing_buys.items() if ref not in proposal_refs),
                  "trades_before": current}
        if not apply:
            return result
        BACKUPS.mkdir(parents=True, exist_ok=True)
        backup = BACKUPS / f"portfolio-pre-reconcile-{datetime.now(timezone.utc):%Y%m%d-%H%M%S-%f}.db"
        with sqlite3.connect(backup) as dest:
            conn.backup(dest)
            if dest.execute("PRAGMA integrity_check").fetchone()[0] != "ok":
                raise RuntimeError("Backup integrity check failed")
        result["backup"] = str(backup)
        conn.execute("PRAGMA foreign_keys=ON")
        conn.execute("BEGIN IMMEDIATE")
        try:
            conn.execute("UPDATE _audit_context SET reason=?", (f"Verified 2024-25 Excel reconciliation {recon_hash[:12]}",))
            inserted = 0
            for row, refs in pending:
                entry, exit_, symbol, sector, buy, sell, _, _, qty, _, pnl, wl, pct, account, _, ref, _ = row
                trade_id = str(uuid.uuid5(uuid.NAMESPACE_URL, f"portfolio-reconciliation|{original_hash}|{ref}"))
                existing = conn.execute("SELECT id FROM trades WHERE id=?", (trade_id,)).fetchone()
                if not existing:
                    duplicate = conn.execute(
                        "SELECT id FROM trades WHERE account_id=? AND symbol=? AND date_exit=? AND ABS(volume-?)<0.000001 AND ABS(price_exit-?)<0.00001",
                        (account, symbol, iso(exit_), float(qty), float(sell)),
                    ).fetchone()
                    if duplicate:
                        raise ValueError(f"Existing matching sale for {ref}: {duplicate[0]}")
                    fx_entry = fx_exit = 1.0
                    if account == "dime":
                        for date, label in ((iso(entry), "entry"), (iso(exit_), "exit")):
                            fx = conn.execute("SELECT rate FROM fx_rates WHERE base='USD' AND quote='THB' AND date<=? ORDER BY date DESC LIMIT 1", (date,)).fetchone()
                            if not fx or not 25 < fx[0] < 45:
                                raise ValueError(f"Historical USD/THB unavailable for {ref} {date}")
                            if label == "entry": fx_entry = fx[0]
                            else: fx_exit = fx[0]
                    conn.execute(
                        """INSERT INTO trades (id,account_id,symbol,resolved_symbol,market,sector,date_entry,date_exit,
                           price_entry,price_exit,volume,amount,pnl_amount,win_loss,pnl_percent,currency,
                           exchange_rate,exit_exchange_rate,note,fee_entry,fee_exit,fee_detail)
                           VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                        (trade_id,account,symbol,symbol if account == "dime" else symbol+".BK",
                         "US" if account == "dime" else "TH",sector,iso(entry),iso(exit_),float(buy),float(sell),float(qty),
                         round(float(buy)*float(qty),8),float(pnl),wl,float(pct),"USD" if account == "dime" else "THB",
                         fx_entry,fx_exit,
                         f"Historical Excel reconciliation; source {ref}; cited buys {', '.join(refs)}; original SHA256 {original_hash}; reconciliation SHA256 {recon_hash}. Source P&L is gross; broker fees unavailable.",
                         None,None,None),
                    )
                    inserted += 1
            # Earlier broker-matched cash offsets had absorbed the missing
            # closed-trade P&L. Release exactly that amount in native currency
            # so importing the history does not alter the observed cash today.
            # These are accounting offsets, not deposits or withdrawals.
            offset_count = 0
            for account, currency in (("dime", "USD"), ("finansia", "THB")):
                amount = -round(sum(float(r[0][10]) for r in pending if r[0][13] == account), 2)
                if not amount:
                    continue
                adj_id = str(uuid.uuid5(uuid.NAMESPACE_URL, f"portfolio-reconciliation-cash-offset|{original_hash}|{account}"))
                if conn.execute("SELECT 1 FROM cash_adjustments WHERE id=?", (adj_id,)).fetchone():
                    continue
                conn.execute(
                    """INSERT INTO cash_adjustments (id,account_id,date,amount,currency,category,note)
                       VALUES (?,?,?,?,?,?,?)""",
                    (adj_id,account,"2026-09-26",amount,currency,"DATA_FIX",
                     f"Releases earlier broker-matched cash offset absorbed by missing 2024-25 closed-trade P&L; source reconciliation SHA256 {recon_hash}; no deposit/withdrawal."),
                )
                offset_count += 1
            for staged_id, ref, decision, reason, row in decisions:
                payload = {"reconciliation_sha256": recon_hash, "original_sha256": original_hash,
                           "source_ref": ref, "confidence_in_workbook": row[14], "reason": reason,
                           "proposed_entry": iso(row[0]), "proposed_exit": iso(row[1]),
                           "proposed_price_entry": row[4], "proposed_price_exit": row[5],
                           "proposed_quantity": row[8], "proposed_pnl": row[10],
                           "missing_buy_proposals": missing_buys.get(ref, []),
                           "trade_id": str(uuid.uuid5(uuid.NAMESPACE_URL, f"portfolio-reconciliation|{original_hash}|{ref}")) if decision == "IMPORTED" else None}
                conn.execute("UPDATE portfolio_history_review SET review_decision=?,review_note=? WHERE id=?",
                             (decision,json.dumps(payload,ensure_ascii=False),staged_id))
            conn.execute("UPDATE _audit_context SET reason=NULL")
            if conn.execute("PRAGMA integrity_check").fetchone()[0] != "ok":
                raise RuntimeError("Post-write integrity check failed")
            conn.commit()
        except Exception:
            conn.rollback()
            raise
        result.update(inserted=inserted, cash_offset_rows_inserted=offset_count,
                      trades_after=conn.execute("SELECT COUNT(*) FROM trades").fetchone()[0],
                      decisions=dict(Counter(r[0] for r in conn.execute("SELECT review_decision FROM portfolio_history_review WHERE review_note LIKE ?", (f'%{recon_hash}%',)))))
        return result


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--reconciliation", dest="recon_path", type=Path, default=RECON)
    parser.add_argument("--original", dest="original_path", type=Path, default=ORIGINAL)
    parser.add_argument("--db", dest="db_path", type=Path, default=DB)
    parser.add_argument("--apply", action="store_true")
    print(json.dumps(run(**vars(parser.parse_args())), ensure_ascii=False, indent=2))
